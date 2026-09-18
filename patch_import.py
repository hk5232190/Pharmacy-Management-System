import re

with open('e:/Projects/PMS-Software/backend/api/v1/medicine.py', 'r', encoding='utf-8') as f:
    content = f.read()

import_statement = """from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import openpyxl
from pydantic import ValidationError"""
if "import openpyxl" not in content:
    content = content.replace("import io\n", f"import io\n{import_statement}\n")

# Remove the old `/import` endpoint and replace with the 3 new endpoints
# The old import endpoint starts at `@router.post("/import", summary="Import medicines from CSV or Excel")`
# and ends right before `@router.put("/{medicine_id}", response_model=BaseResponse[MedicineResponse], summary="Update a medicine")`

old_endpoint_regex = re.compile(
    r'@router\.post\("/import", summary="Import medicines from CSV or Excel"\).*?(?=@router\.put\("/{medicine_id}")',
    re.DOTALL
)

new_endpoints = '''@router.get("/import-template", summary="Download smart Excel template for import")
def download_import_template(db: Session = Depends(get_db)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicines"

    headers = [
        "BrandName*", "GenericName*", "Category*", "Company*", 
        "Unit", "DosageForm", "Strength", "ReorderLevel",
        "RackNumber", "DefaultCostPrice", "DefaultSellingPrice", "Barcode", "RequiresPrescription(0/1)"
    ]
    
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
        ws.column_dimensions[get_column_letter(col_num)].width = 20
        
    categories = db.query(Category).all()
    companies = db.query(Company).all()
    
    cat_names = [c.CategoryName.replace(",", "") for c in categories if c.CategoryName]
    comp_names = [c.CompanyName.replace(",", "") for c in companies if c.CompanyName]
    
    if cat_names:
        # Excel data validation formula1 limit is 255 chars. 
        # If it's too long, we might need a hidden sheet, but for simplicity we'll try direct list
        cat_formula = '"' + ",".join(cat_names)[:253] + '"'
        dv_cat = DataValidation(type="list", formula1=cat_formula, allow_blank=True)
        ws.add_data_validation(dv_cat)
        dv_cat.add('C2:C1000')
        
    if comp_names:
        comp_formula = '"' + ",".join(comp_names)[:253] + '"'
        dv_comp = DataValidation(type="list", formula1=comp_formula, allow_blank=True)
        ws.add_data_validation(dv_comp)
        dv_comp.add('D2:D1000')
        
    # Unit dropdown
    dv_unit = DataValidation(type="list", formula1='"Box,Strip,Bottle,Tube,Injection,Pieces"', allow_blank=True)
    ws.add_data_validation(dv_unit)
    dv_unit.add('E2:E1000')

    # Dosage dropdown
    dv_dosage = DataValidation(type="list", formula1='"Tablet,Capsule,Syrup,Injection,Cream,Drops,Ointment,Other"', allow_blank=True)
    ws.add_data_validation(dv_dosage)
    dv_dosage.add('F2:F1000')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=medicines_import_template.xlsx"}
    )

@router.post("/import-preview", summary="Preview medicines import and map categories/companies")
def preview_medicines_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith(('.csv', '.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Only CSV and Excel files are allowed")
        
    contents = file.file.read()
    rows = []
    
    if file.filename.endswith('.csv'):
        decoded = contents.decode('utf-8')
        csv_reader = csv.DictReader(io.StringIO(decoded))
        rows = list(csv_reader)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        sheet_rows = list(sheet.iter_rows(values_only=True))
        if sheet_rows:
            headers = [str(cell).strip() if cell is not None else "" for cell in sheet_rows[0]]
            # normalize headers
            headers = [h.replace("*", "").replace("(0/1)", "").strip() for h in headers]
            for row in sheet_rows[1:]:
                if any(cell is not None and str(cell).strip() for cell in row):
                    row_dict = {headers[i]: str(cell).strip() if cell is not None else "" for i, cell in enumerate(row) if i < len(headers)}
                    rows.append(row_dict)
                    
    # Pre-fetch for quick mapping
    categories = {c.CategoryName.lower(): c.CategoryId for c in db.query(Category).all() if c.CategoryName}
    companies = {c.CompanyName.lower(): c.CompanyId for c in db.query(Company).all() if c.CompanyName}
    
    preview_data = []
    for i, row in enumerate(rows, start=1):
        row_preview = {
            "RowNumber": i,
            "BrandName": row.get("BrandName", ""),
            "GenericName": row.get("GenericName", ""),
            "CategoryName": row.get("Category", row.get("CategoryName", "")),
            "CompanyName": row.get("Company", row.get("CompanyName", "")),
            "Unit": row.get("Unit", "Box"),
            "DosageForm": row.get("DosageForm", ""),
            "Strength": row.get("Strength", ""),
            "ReorderLevel": row.get("ReorderLevel", 10),
            "RackNumber": row.get("RackNumber", ""),
            "DefaultCostPrice": row.get("DefaultCostPrice", 0),
            "DefaultSellingPrice": row.get("DefaultSellingPrice", 0),
            "Barcode": row.get("Barcode", ""),
            "RequiresPrescription": row.get("RequiresPrescription", False),
            
            "CategoryId": None,
            "CompanyId": None,
            "IsValid": True,
            "Errors": []
        }
        
        if not row_preview["BrandName"]:
            row_preview["IsValid"] = False
            row_preview["Errors"].append("Brand Name is required")
        if not row_preview["GenericName"]:
            row_preview["IsValid"] = False
            row_preview["Errors"].append("Generic Name is required")
            
        cat_name = row_preview["CategoryName"].lower() if row_preview["CategoryName"] else ""
        comp_name = row_preview["CompanyName"].lower() if row_preview["CompanyName"] else ""
        
        if cat_name in categories:
            row_preview["CategoryId"] = categories[cat_name]
        else:
            row_preview["IsValid"] = False
            row_preview["Errors"].append(f"Category '{row_preview['CategoryName']}' not found in DB")
            
        if comp_name in companies:
            row_preview["CompanyId"] = companies[comp_name]
        else:
            row_preview["IsValid"] = False
            row_preview["Errors"].append(f"Company '{row_preview['CompanyName']}' not found in DB")
            
        # Optional defaults mapping
        try:
            row_preview["ReorderLevel"] = int(float(row_preview["ReorderLevel"]) if str(row_preview["ReorderLevel"]).strip() else 10)
        except ValueError:
            row_preview["ReorderLevel"] = 10
            
        try:
            row_preview["DefaultCostPrice"] = float(row_preview["DefaultCostPrice"]) if str(row_preview["DefaultCostPrice"]).strip() else 0.0
        except ValueError:
            row_preview["DefaultCostPrice"] = 0.0
            
        try:
            row_preview["DefaultSellingPrice"] = float(row_preview["DefaultSellingPrice"]) if str(row_preview["DefaultSellingPrice"]).strip() else 0.0
        except ValueError:
            row_preview["DefaultSellingPrice"] = 0.0

        if str(row_preview["RequiresPrescription"]) in ['1', 'true', 'True', 'yes', 'Yes']:
            row_preview["RequiresPrescription"] = True
        else:
            row_preview["RequiresPrescription"] = False

        preview_data.append(row_preview)
        
    return {"success": True, "data": preview_data, "message": "Preview generated"}

@router.post("/import-bulk", summary="Bulk create validated medicines")
def bulk_import_medicines(
    medicines: List[MedicineCreate],
    db: Session = Depends(get_db)
):
    count = 0
    for med_in in medicines:
        new_med = Medicine(**med_in.model_dump())
        db.add(new_med)
        count += 1
    
    db.commit()
    return {"success": True, "message": f"Successfully imported {count} medicines"}

'''

content = old_endpoint_regex.sub(new_endpoints, content)

with open('e:/Projects/PMS-Software/backend/api/v1/medicine.py', 'w', encoding='utf-8') as f:
    f.write(content)
