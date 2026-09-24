from fastapi import APIRouter, Depends, HTTPException, Query, File, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from sqlalchemy import or_, func
from typing import List
import random
import csv
import io
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import openpyxl
from pydantic import ValidationError
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import date, datetime
from typing import Optional

from models import Medicine, Category, Company, StockBatch, SaleItem, PurchaseItem, InventorySettings, OpeningStockEntry, OpeningStockItem, StockAdjustment, AuditLog
from schemas.medicine import MedicineCreate, MedicineUpdate, MedicineResponse, InitialStockBatch
from schemas.base import BaseResponse
from schemas.opening_stock import normalize_batch_code, compute_payload_hash, OpeningStockLineItem
from api.deps import get_current_user, get_db
from core.logger import logger

router = APIRouter(dependencies=[Depends(get_current_user)])


# ─────────────────────────────────────────────────────────────────────────────
# Internal helpers
# ─────────────────────────────────────────────────────────────────────────────

def _generate_initial_stock_reference(db: Session) -> str:
    """Generate a unique IS-YYYYMMDD-NNN reference for initial stock sessions."""
    today_str = datetime.utcnow().strftime("%Y%m%d")
    prefix = f"IS-{today_str}-"
    existing = (
        db.query(OpeningStockEntry)
        .filter(OpeningStockEntry.ReferenceNo.like(f"{prefix}%"))
        .count()
    )
    seq = existing + 1
    for _ in range(100):
        candidate = f"{prefix}{seq:03d}"
        if not db.query(OpeningStockEntry).filter(OpeningStockEntry.ReferenceNo == candidate).first():
            return candidate
        seq += 1
    raise RuntimeError("Could not generate a unique reference number")


def _commit_initial_stock_batches(
    db: Session,
    medicine_id: int,
    batches: List[InitialStockBatch],
    user_id: int,
    medicine_name: str,
) -> str:
    """
    Atomically create StockBatch + OpeningStockEntry/Item + StockAdjustment records
    for the given batches, reusing Opening Stock logic.

    Returns the reference number of the created session.
    Must be called inside an already-open transaction; caller handles commit/rollback.
    """
    # Build canonical line items to generate the idempotency token
    line_items = [
        OpeningStockLineItem(
            MedicineId=medicine_id,
            BatchCode=b.BatchCode,
            Quantity=b.Quantity,
            CostPrice=b.CostPrice,
            SellingPrice=b.SellingPrice,
            ExpiryDate=b.ExpiryDate,
            ManufacturingDate=b.ManufacturingDate,
        )
        for b in batches
    ]

    # Duplicate batch code check within this request
    seen_codes: set = set()
    for batch in batches:
        code = normalize_batch_code(batch.BatchCode)
        if code in seen_codes:
            raise HTTPException(
                status_code=422,
                detail=f"Duplicate batch '{code}' in initial stock — each batch must be unique."
            )
        seen_codes.add(code)
        # Check DB conflict
        existing = (
            db.query(StockBatch)
            .filter(
                StockBatch.MedicineId == medicine_id,
                func.upper(func.trim(StockBatch.BatchCode)) == code,
            )
            .first()
        )
        if existing:
            raise HTTPException(
                status_code=409,
                detail=f"Batch '{code}' already exists for '{medicine_name}'."
            )

    token = compute_payload_hash(line_items)
    reference_no = _generate_initial_stock_reference(db)
    total_value = sum(Decimal(str(b.Quantity)) * b.CostPrice for b in batches)

    entry = OpeningStockEntry(
        ReferenceNo=reference_no,
        ImportHash=token,
        Notes=f"Initial stock added with medicine '{medicine_name}'",
        CreatedBy=user_id,
        TotalItems=len(batches),
        TotalValue=total_value,
        Status="ACTIVE",
    )
    db.add(entry)
    db.flush()  # get EntryId

    for batch in batches:
        code = normalize_batch_code(batch.BatchCode)

        new_batch = StockBatch(
            MedicineId=medicine_id,
            BatchCode=code,
            Quantity=batch.Quantity,
            CostPrice=batch.CostPrice,
            SellingPrice=batch.SellingPrice,
            ManufacturingDate=batch.ManufacturingDate,
            ExpiryDate=batch.ExpiryDate,
            Source="OPENING_STOCK",
        )
        db.add(new_batch)
        db.flush()  # get BatchId

        os_item = OpeningStockItem(
            EntryId=entry.EntryId,
            BatchId=new_batch.BatchId,
            MedicineId=medicine_id,
            BatchCode=code,
            Quantity=batch.Quantity,
            CostPrice=batch.CostPrice,
            SellingPrice=batch.SellingPrice,
            ExpiryDate=batch.ExpiryDate,
            ManufacturingDate=batch.ManufacturingDate,
        )
        db.add(os_item)

        # Stock movement record — Reason prefix "OPENING_STOCK:" is recognized
        # by get_stock_movements() as "Opening Stock" type, not "Stock Adjustment".
        adj = StockAdjustment(
            BatchId=new_batch.BatchId,
            UserId=user_id,
            AdjustmentType="Increase",
            Quantity=batch.Quantity,
            PreviousQuantity=0,
            NewQuantity=batch.Quantity,
            Reason=f"OPENING_STOCK:{reference_no}",
        )
        db.add(adj)

    db.add(
        AuditLog(
            UserId=user_id,
            Action="INITIAL_STOCK_ENTRY",
            Description=(
                f"Initial Stock session {reference_no} committed with medicine '{medicine_name}': "
                f"{len(batches)} batch(es), total value {total_value}."
            ),
        )
    )

    return reference_no


# ─────────────────────────────────────────────────────────────────────────────
# List medicines
# ─────────────────────────────────────────────────────────────────────────────

@router.get("", summary="Get all medicines")
def get_medicines(
    search: str = Query(None, description="Search by name, generic name, or barcode"),
    category_id: int = Query(None, description="Filter by CategoryId"),
    company_id: int = Query(None, description="Filter by CompanyId"),
    page: int = Query(1, ge=1, description="Page number"),
    page_size: int = Query(25, ge=0, description="Items per page. 0 for all."),
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    query = db.query(Medicine, Category.CategoryName, Company.CompanyName)\
              .outerjoin(Category, Medicine.CategoryId == Category.CategoryId)\
              .outerjoin(Company, Medicine.CompanyId == Company.CompanyId)
    
    if search:
        search_term = search.strip()
        search_id = None
        if search_term.upper().startswith("MED-"):
            id_part = search_term[4:].lstrip("0")
            if id_part.isdigit():
                search_id = int(id_part)
        elif search_term.isdigit():
            search_id = int(search_term)
            
        conditions = [
            Medicine.BrandName.ilike(f"%{search_term}%"),
            Medicine.GenericName.ilike(f"%{search_term}%"),
            Medicine.Barcode.ilike(f"%{search_term}%")
        ]
        if search_id is not None:
            conditions.append(Medicine.MedicineId == search_id)
            
        query = query.filter(or_(*conditions))
    if category_id:
        query = query.filter(Medicine.CategoryId == category_id)
    if company_id:
        query = query.filter(Medicine.CompanyId == company_id)
        
    total = query.count()
    
    if page_size > 0:
        query = query.order_by(Medicine.BrandName).offset((page - 1) * page_size).limit(page_size)
    else:
        query = query.order_by(Medicine.BrandName)
        
    results = query.all()
    
    medicines_list = []
    for med, cat_name, comp_name in results:
        med_dict = {c.name: getattr(med, c.name) for c in med.__table__.columns}
        med_dict["CategoryName"] = cat_name
        med_dict["CompanyName"] = comp_name
        medicines_list.append(med_dict)
        
    return {"success": True, "data": medicines_list, "total": total, "page": page, "page_size": page_size}


# ─────────────────────────────────────────────────────────────────────────────
# Create medicine (with optional Initial Stock)
# ─────────────────────────────────────────────────────────────────────────────

@router.post("", response_model=BaseResponse[MedicineResponse], summary="Create a new medicine")
def create_medicine(
    medicine_in: MedicineCreate,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    # Fetch Inventory Settings for Auto-Barcode preference
    inv_settings = db.query(InventorySettings).first()
    auto_generate = inv_settings.AutoGenerateBarcode if inv_settings else True

    # Auto-generate EAN-13 barcode if none provided and setting is enabled
    if not medicine_in.Barcode and auto_generate:
        while True:
            # Generate 12 random digits for EAN-13
            base = str(random.randint(100000000000, 999999999999))
            
            # Calculate EAN-13 check digit
            total = 0
            for i, char in enumerate(base):
                if i % 2 == 0:
                    total += int(char) * 1
                else:
                    total += int(char) * 3
            check_digit = (10 - (total % 10)) % 10
            
            generated_barcode = base + str(check_digit)
            
            # Uniqueness check
            existing = db.query(Medicine).filter(Medicine.Barcode == generated_barcode).first()
            if not existing:
                medicine_in.Barcode = generated_barcode
                break

    # Optional: check if barcode already exists
    if medicine_in.Barcode:
        existing = db.query(Medicine).filter(Medicine.Barcode == medicine_in.Barcode).first()
        if existing:
            raise HTTPException(status_code=400, detail="Medicine with this barcode already exists")
            
    # Apply defaults from Inventory Settings
    if not medicine_in.Unit:
        medicine_in.Unit = inv_settings.DefaultUnit if inv_settings else "Box"
    
    if medicine_in.ReorderLevel is None or medicine_in.ReorderLevel == 0:
        medicine_in.ReorderLevel = inv_settings.LowStockThreshold if inv_settings else 10
        
    # Calculate precision selling price if not provided
    if medicine_in.DefaultCostPrice and (not medicine_in.DefaultSellingPrice or medicine_in.DefaultSellingPrice == 0):
        margin = Decimal(str(inv_settings.DefaultProfitMargin)) if inv_settings else Decimal('0')
        cost = Decimal(str(medicine_in.DefaultCostPrice))
        selling = cost * (Decimal('1') + margin / Decimal('100'))
        medicine_in.DefaultSellingPrice = float(selling.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))

    # Extract initial_stock before dumping medicine fields
    initial_stock = medicine_in.initial_stock

    try:
        med_data = medicine_in.model_dump(exclude={"initial_stock", "ExistingMedicineId"})
        new_medicine = Medicine(**med_data)
        db.add(new_medicine)
        db.flush()  # get MedicineId before creating stock batches

        initial_stock_ref = None
        if initial_stock:
            initial_stock_ref = _commit_initial_stock_batches(
                db=db,
                medicine_id=new_medicine.MedicineId,
                batches=initial_stock,
                user_id=current_user.UserId,
                medicine_name=new_medicine.BrandName,
            )

        db.commit()
        db.refresh(new_medicine)
        
        # Load relationships for response
        med = db.query(Medicine).filter(Medicine.MedicineId == new_medicine.MedicineId).first()
        response_data = {c.name: getattr(med, c.name) for c in med.__table__.columns}
        response_data["CategoryName"] = med.category.CategoryName if med.category else None
        response_data["CompanyName"] = med.company.CompanyName if med.company else None
        
        msg = "Medicine created successfully"
        if initial_stock_ref:
            msg += f" with initial stock (session {initial_stock_ref})"
        
        return {"data": response_data, "message": msg}

    except HTTPException:
        db.rollback()
        raise
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="Database integrity error — check for duplicate barcode or batch code.")
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


# ─────────────────────────────────────────────────────────────────────────────
# Export medicines — CSV (medicine details only, for quick reference)
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/export", summary="Export all medicines to CSV (medicine details only)")
def export_medicines(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    from sqlalchemy.orm import joinedload
    medicines = db.query(Medicine).options(joinedload(Medicine.category), joinedload(Medicine.company)).all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow([
        "Brand Name", "Formula", "Category", "Company",
        "Unit", "Dosage Form", "Reorder Level (Min Stock)",
        "Rack Number", "Status"
    ])
    
    for med in medicines:
        cat_name = med.category.CategoryName if med.category else ""
        comp_name = med.company.CompanyName if med.company else ""
        status = "Active" if getattr(med, "IsActive", True) else "Inactive"
        writer.writerow([
            med.BrandName or "", med.GenericName or "",
            cat_name, comp_name,
            med.Unit or "", med.DosageForm or "",
            med.ReorderLevel, med.RackNumber or "", status
        ])
        
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=medicines_export.csv"}
    )


# ─────────────────────────────────────────────────────────────────────────────
# Export with Stock — XLSX (medicine + current stock batches, re-importable)
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/export-with-stock", summary="Export medicines with current stock batches as re-importable XLSX")
def export_medicines_with_stock(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    from sqlalchemy.orm import joinedload
    from openpyxl.styles import Font, PatternFill, Alignment

    medicines = (
        db.query(Medicine)
        .options(
            joinedload(Medicine.category),
            joinedload(Medicine.company),
            joinedload(Medicine.batches),
        )
        .order_by(Medicine.BrandName)
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicines"

    med_headers = [
        "Brand Name*", "Formula*", "Category*", "Company*",
        "Unit*", "Dosage Form*", "Reorder Level (Min Stock)", "Rack Number", "Status*"
    ]
    stock_headers = [
        "Batch Number (Optional)", "Quantity (Optional)", "Cost Price (Optional)",
        "Selling Price (Optional)", "Expiry Date YYYY-MM-DD (Optional)", "Mfg Date YYYY-MM-DD (Optional)"
    ]
    all_headers = med_headers + stock_headers

    med_fill  = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    stock_fill = PatternFill(start_color="1A5C3A", end_color="1A5C3A", fill_type="solid")
    white_bold = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.fill = med_fill if col <= len(med_headers) else stock_fill
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws.cell(row=2, column=1, value="← Medicine Details")
    ws.cell(row=2, column=len(med_headers) + 1, value="← Initial Stock (optional)")
    for col in range(1, len(all_headers) + 1):
        ws.cell(row=2, column=col).font = Font(italic=True, color="888888")
    ws.freeze_panes = "A3"

    data_row = 3
    for med in medicines:
        cat_name  = med.category.CategoryName if med.category else ""
        comp_name = med.company.CompanyName   if med.company  else ""
        status    = "Active" if med.IsActive else "Inactive"
        med_row   = [
            med.BrandName or "", med.GenericName or "",
            cat_name, comp_name,
            med.Unit or "", med.DosageForm or "",
            med.ReorderLevel, med.RackNumber or "", status,
        ]

        # Only include batches with stock > 0
        active_batches = [b for b in (med.batches or []) if b.Quantity > 0]

        if active_batches:
            for batch in active_batches:
                row_data = med_row + [
                    batch.BatchCode,
                    batch.Quantity,
                    float(batch.CostPrice),
                    float(batch.SellingPrice),
                    batch.ExpiryDate.strftime("%Y-%m-%d") if batch.ExpiryDate else "",
                    batch.ManufacturingDate.strftime("%Y-%m-%d") if batch.ManufacturingDate else "",
                ]
                for col, val in enumerate(row_data, 1):
                    ws.cell(row=data_row, column=col, value=val)
                data_row += 1
        else:
            # Medicine with no stock — write one row, leave batch columns empty
            for col, val in enumerate(med_row + [""] * 6, 1):
                ws.cell(row=data_row, column=col, value=val)
            data_row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=medicines_with_stock_export.xlsx"}
    )


# ─────────────────────────────────────────────────────────────────────────────
# Download import template (Excel)
# ─────────────────────────────────────────────────────────────────────────────

@router.get("/import-template", summary="Download smart Excel import template with Instructions sheet")
def download_import_template(db: Session = Depends(get_db)):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()

    # ── Medicines data sheet ───────────────────────────────────────────────
    ws = wb.active
    ws.title = "Medicines"

    med_headers = [
        "Brand Name*", "Formula*", "Category*", "Company*",
        "Unit*", "Dosage Form*", "Reorder Level (Min Stock)", "Rack Number", "Status*"
    ]
    stock_headers = [
        "Batch Number (Optional)", "Quantity (Optional)", "Cost Price (Optional)",
        "Selling Price (Optional)", "Expiry Date YYYY-MM-DD (Optional)", "Mfg Date YYYY-MM-DD (Optional)"
    ]
    all_headers = med_headers + stock_headers
    total_cols   = len(all_headers)

    med_fill   = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    stock_fill = PatternFill(start_color="1A5C3A", end_color="1A5C3A", fill_type="solid")
    white_bold = Font(color="FFFFFF", bold=True)
    example_font = Font(color="888888", italic=True)
    example_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

    # Row 1 — column headers
    for col, header in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.fill = med_fill if col <= len(med_headers) else stock_fill
        ws.column_dimensions[get_column_letter(col)].width = 23
    ws.row_dimensions[1].height = 30

    # Row 2 — section labels
    ws.cell(row=2, column=1, value="← Medicine Details (required)")
    ws.cell(row=2, column=len(med_headers)+1, value="← Initial Stock (optional — leave blank for zero-stock import)")
    for col in range(1, total_cols+1):
        ws.cell(row=2, column=col).font = Font(italic=True, color="888888")

    # Dropdowns (apply from row 3 onward)
    categories = db.query(Category).all()
    companies  = db.query(Company).all()
    cat_names  = [c.CategoryName.replace(",", "") for c in categories if c.CategoryName]
    comp_names = [c.CompanyName.replace(",", "")  for c in companies  if c.CompanyName]

    if cat_names:
        dv = DataValidation(type="list", formula1='"' + ",".join(cat_names)[:253] + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add("C3:C2000")
    if comp_names:
        dv = DataValidation(type="list", formula1='"' + ",".join(comp_names)[:253] + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add("D3:D2000")

    dv_unit = DataValidation(type="list", formula1='"Box,Strip,Bottle,Tube,Piece,Vial,Ampoule,Sachet,Pack,Jar,Can"', allow_blank=True)
    ws.add_data_validation(dv_unit); dv_unit.add("E3:E2000")

    dv_dosage = DataValidation(type="list", formula1='"Tablet,Capsule,Syrup,Suspension,Injection,Cream,Ointment,Drops,Gel,Lotion,Spray,Inhaler,Powder,Suppository,Other"', allow_blank=True)
    ws.add_data_validation(dv_dosage); dv_dosage.add("F3:F2000")

    dv_status = DataValidation(type="list", formula1='"Active,Inactive"', allow_blank=True)
    ws.add_data_validation(dv_status); dv_status.add("I3:I2000")

    ws.freeze_panes = "A3"

    # ── Instructions sheet ─────────────────────────────────────────────────
    ws_i = wb.create_sheet("Instructions")
    ws_i.column_dimensions["A"].width = 90
    title_font  = Font(bold=True, size=13)
    head_font   = Font(bold=True, size=11)
    body_font   = Font(size=10)
    code_font   = Font(name="Courier New", size=9, color="1A5C3A")
    head_fill   = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    def _irow(row, value, fnt=None, fill=None):
        cell = ws_i.cell(row=row, column=1, value=value)
        if fnt:  cell.font = fnt
        if fill: cell.fill = fill
        ws_i.row_dimensions[row].height = 18
        return row + 1

    r = 1
    r = _irow(r, "Medicine Import Template — Instructions", title_font, head_fill)
    r = _irow(r, "")
    r = _irow(r, "COLUMN GUIDE", head_font)
    r = _irow(r, "Columns A–I  (Blue)  — Medicine Details. All marked * are REQUIRED.", body_font)
    r = _irow(r, "Columns J–O  (Green) — Initial Stock Batch. Leave ALL empty to import medicine with zero stock.", body_font)
    r = _irow(r, "")
    r = _irow(r, "SCENARIO 1 — Medicine + Initial Stock", head_font)
    r = _irow(r, "Fill medicine details (A–I) and batch details (J–O). The system will create the medicine and stock atomically.", body_font)
    r = _irow(r, "Affects: Current Stock, Stock Movement History, Expiry Tracking, Dashboard, POS/FEFO.", body_font)
    r = _irow(r, "")
    r = _irow(r, "SCENARIO 2 — Medicine Only (no stock)", head_font)
    r = _irow(r, "Fill medicine details (A–I) only. Leave columns J–O completely blank.", body_font)
    r = _irow(r, "The medicine is imported with zero stock. Add stock later via Purchases or Opening Stock.", body_font)
    r = _irow(r, "")
    r = _irow(r, "MULTIPLE BATCHES FOR THE SAME MEDICINE", head_font)
    r = _irow(r, "Repeat the medicine row with the SAME Brand Name and different batch data:", body_font)
    r = _irow(r, "  Panadol | Paracetamol | ... | BT-001 | 100 | 50 | 75 | 2026-12-31 |", code_font)
    r = _irow(r, "  Panadol | Paracetamol | ... | BT-002 | 200 | 50 | 75 | 2027-06-30 |", code_font)
    r = _irow(r, "Both rows create ONE medicine with TWO stock batches.", body_font)
    r = _irow(r, "")
    r = _irow(r, "ADDING STOCK TO AN EXISTING MEDICINE", head_font)
    r = _irow(r, "If a medicine with the same Brand Name already exists in the system, the import will", body_font)
    r = _irow(r, "skip creating a duplicate and ONLY add the new batch(es) to the existing medicine record.", body_font)
    r = _irow(r, "")
    r = _irow(r, "DATE FORMAT", head_font)
    r = _irow(r, "Use YYYY-MM-DD format for Expiry Date and Manufacturing Date (e.g. 2026-12-31).", body_font)
    r = _irow(r, "")
    r = _irow(r, "CATEGORY / COMPANY", head_font)
    r = _irow(r, "Must exactly match an existing Category/Company name. Use the dropdown in column C/D.", body_font)
    r = _irow(r, "Unknown values can be fixed in the Import Preview screen before saving.", body_font)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=medicines_import_template.xlsx"}
    )


# ─────────────────────────────────────────────────────────────────────────────
# Row field extractor helper
# ─────────────────────────────────────────────────────────────────────────────

def _row_get(row: dict, *keys: str) -> str:
    """Return the first non-empty value from a dict row matching any of the given key prefixes."""
    for key in keys:
        key_lower = key.lower()
        for header, v in row.items():
            if header.strip().lower().startswith(key_lower):
                val = str(v).strip() if v is not None else ""
                if val and val.lower() not in ("none", "nan"):
                    return val
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# Parse initial stock from a single CSV/XLSX row
# ─────────────────────────────────────────────────────────────────────────────

def _parse_initial_stock_from_row(row: dict, row_num: int) -> tuple[Optional[dict], Optional[str]]:
    """
    Parse the optional initial stock fields from a CSV/XLSX import row.
    Returns (batch_dict, error_str). batch_dict is None if no stock columns present.
    """
    batch_code  = _row_get(row, "batch number", "batchcode", "batch no")
    qty_raw     = _row_get(row, "quantity")
    cost_raw    = _row_get(row, "cost price", "costprice")
    selling_raw = _row_get(row, "selling price", "sellingprice")
    expiry_raw  = _row_get(row, "expiry date", "expirydate")
    mfg_raw     = _row_get(row, "mfg date", "manufacturing date", "manufacturingdate")

    if not any([batch_code, qty_raw, cost_raw, selling_raw, expiry_raw]):
        return None, None  # No stock data in this row

    errors: list[str] = []
    if not batch_code:  errors.append("Batch Number is required when providing initial stock")
    if not qty_raw:     errors.append("Quantity is required when providing initial stock")
    else:
        try:
            qty = int(float(qty_raw))
            if qty <= 0: errors.append("Quantity must be > 0")
        except ValueError:
            errors.append(f"Invalid Quantity '{qty_raw}'"); qty = 0
    if not cost_raw:    errors.append("Cost Price is required when providing initial stock")
    if not selling_raw: errors.append("Selling Price is required when providing initial stock")
    if not expiry_raw:  errors.append("Expiry Date is required when providing initial stock")
    if errors:
        return None, "; ".join(errors)

    try:
        qty     = int(float(qty_raw))
        cost    = float(Decimal(str(cost_raw).replace(",", "").strip()))
        selling = float(Decimal(str(selling_raw).replace(",", "").strip()))
    except (ValueError, InvalidOperation) as e:
        return None, f"Invalid numeric value: {e}"

    expiry_date = None
    for parse_fn in [lambda s: date.fromisoformat(s[:10]), lambda s: datetime.strptime(s, "%d/%m/%Y").date()]:
        try:
            expiry_date = parse_fn(expiry_raw.strip()); break
        except ValueError:
            continue
    if not expiry_date:
        return None, f"Invalid Expiry Date '{expiry_raw}'. Use YYYY-MM-DD."

    mfg_date = None
    if mfg_raw and mfg_raw.lower() not in ("", "none", "n/a", "-"):
        for parse_fn in [lambda s: date.fromisoformat(s[:10]), lambda s: datetime.strptime(s, "%d/%m/%Y").date()]:
            try:
                mfg_date = parse_fn(mfg_raw.strip()); break
            except ValueError:
                continue
        if not mfg_date:
            return None, f"Invalid Manufacturing Date '{mfg_raw}'. Use YYYY-MM-DD."
        if mfg_date > expiry_date:
            return None, "Manufacturing Date must be on or before Expiry Date."

    return {
        "BatchCode": batch_code.strip().upper(),
        "Quantity": qty,
        "CostPrice": cost,
        "SellingPrice": selling,
        "ExpiryDate": expiry_date.isoformat(),
        "ManufacturingDate": mfg_date.isoformat() if mfg_date else None,
    }, None


# ─────────────────────────────────────────────────────────────────────────────
# Preview import — grouped by medicine (BrandName), multi-batch aware
# ─────────────────────────────────────────────────────────────────────────────

@router.post("/import-preview", summary="Preview medicines import grouped by medicine")
def preview_medicines_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith(('.csv', '.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Only CSV and Excel files are allowed")

    contents = file.file.read()
    rows: list[dict] = []

    if file.filename.endswith('.csv'):
        decoded = contents.decode('utf-8-sig')  # handle BOM
        rows = list(csv.DictReader(io.StringIO(decoded)))
    else:
        wb_src = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        # Prefer "Medicines" sheet if present (matches template)
        sheet = wb_src["Medicines"] if "Medicines" in wb_src.sheetnames else wb_src.active
        sheet_rows = list(sheet.iter_rows(values_only=True))
        if sheet_rows:
            headers = [
                h.replace("*", "").replace("(0/1)", "").strip()
                for h in (str(c).strip() if c is not None else "" for c in sheet_rows[0])
            ]
            # Detect non-data rows to skip (section labels starting with ←, example rows)
            data_start = 1
            for look in range(1, min(6, len(sheet_rows))):
                row_vals = [str(v).strip() if v else "" for v in sheet_rows[look]]
                first_cell = row_vals[0] if row_vals else ""
                # Skip section label rows and example rows
                if any(v.startswith("\u2190") for v in row_vals) or "example" in first_cell.lower() or "delete" in first_cell.lower():
                    data_start = look + 1
            for row in sheet_rows[data_start:]:
                if any(c is not None and str(c).strip() for c in row):
                    rows.append({headers[i]: str(c).strip() if c is not None else "" for i, c in enumerate(row) if i < len(headers)})

    # Pre-fetch lookups
    cat_map  = {c.CategoryName.lower(): c.CategoryId for c in db.query(Category).all() if c.CategoryName}
    comp_map = {c.CompanyName.lower():  c.CompanyId  for c in db.query(Company).all()  if c.CompanyName}

    # ── Group rows by BrandName (case-insensitive) ─────────────────────────
    # One preview entry per unique medicine; multiple CSV rows for same brand
    # accumulate their batches into InitialBatches.
    from collections import OrderedDict
    groups: "OrderedDict[str, dict]" = OrderedDict()

    for i, row in enumerate(rows, start=1):
        brand = _row_get(row, "brand name", "brandname")
        generic = _row_get(row, "formula", "genericname")

        # Rows without a brand name are invalid stand-alone entries
        group_key = brand.lower() if brand else f"__unnamed_row_{i}__"
        is_new    = group_key not in groups

        if is_new:
            cat_name  = _row_get(row, "category")
            comp_name = _row_get(row, "company")
            status_v  = _row_get(row, "status", "isactive")
            is_active = status_v.lower() in ("active", "true", "1") if status_v else True

            rl_raw = _row_get(row, "reorder level", "reorderlevel")
            try:
                reorder_level = int(float(rl_raw)) if rl_raw else 10
            except (ValueError, TypeError):
                reorder_level = 10

            entry: dict = {
                "RowNumbers":        [i],
                "BrandName":         brand,
                "GenericName":       generic,
                "CategoryName":      cat_name,
                "CompanyName":       comp_name,
                "Unit":              _row_get(row, "unit") or "Box",
                "DosageForm":        _row_get(row, "dosage form", "dosageform"),
                "ReorderLevel":      reorder_level,
                "RackNumber":        _row_get(row, "rack number", "racknumber"),
                "RequiresPrescription": False,
                "IsActive":          is_active,
                "DefaultCostPrice":  0.0,
                "DefaultSellingPrice": 0.0,
                "Barcode":           None,
                "CategoryId":        None,
                "CompanyId":         None,
                "IsValid":           True,
                "Errors":            [],
                "InitialBatches":    [],
                "ExistingMedicineId": None,
            }

            # Validate required medicine fields
            if not brand:   entry["IsValid"] = False; entry["Errors"].append("Brand Name is required")
            if not generic: entry["IsValid"] = False; entry["Errors"].append("Formula is required")

            # Resolve Category
            cat_lower = cat_name.lower()
            if cat_lower in cat_map:
                entry["CategoryId"] = cat_map[cat_lower]
            elif cat_name:
                entry["IsValid"] = False; entry["Errors"].append(f"Category '{cat_name}' not found")
            else:
                entry["IsValid"] = False; entry["Errors"].append("Category is required")

            # Resolve Company
            comp_lower = comp_name.lower()
            if comp_lower in comp_map:
                entry["CompanyId"] = comp_map[comp_lower]
            elif comp_name:
                entry["IsValid"] = False; entry["Errors"].append(f"Company '{comp_name}' not found")
            else:
                entry["IsValid"] = False; entry["Errors"].append("Company is required")

            # Check if medicine already exists in the database
            if brand:
                existing_med = db.query(Medicine).filter(
                    func.lower(Medicine.BrandName) == brand.lower()
                ).first()
                if existing_med:
                    entry["ExistingMedicineId"] = existing_med.MedicineId
                    # Use existing IDs; remove resolution errors since medicine exists
                    entry["CategoryId"] = entry["CategoryId"] or existing_med.CategoryId
                    entry["CompanyId"]  = entry["CompanyId"]  or existing_med.CompanyId
                    entry["Errors"] = [
                        e for e in entry["Errors"]
                        if "Category" not in e and "Company" not in e
                    ]
                    entry["IsValid"] = len(entry["Errors"]) == 0

            groups[group_key] = entry

        else:
            # Continuation row for the same medicine — add row number only
            groups[group_key]["RowNumbers"].append(i)

        # ── Parse initial stock for this row ───────────────────────────────
        entry = groups[group_key]
        batch_dict, stock_err = _parse_initial_stock_from_row(row, i)
        if stock_err:
            entry["IsValid"] = False
            entry["Errors"].append(f"Row {i} batch: {stock_err}")
        elif batch_dict:
            # Duplicate batch code within this medicine group?
            existing_codes = {b["BatchCode"] for b in entry["InitialBatches"]}
            if batch_dict["BatchCode"] in existing_codes:
                entry["IsValid"] = False
                entry["Errors"].append(f"Row {i}: Duplicate Batch '{batch_dict['BatchCode']}' for this medicine")
            else:
                # For existing medicines — check if batch already in DB
                med_id = entry["ExistingMedicineId"]
                if med_id:
                    db_batch = db.query(StockBatch).filter(
                        StockBatch.MedicineId == med_id,
                        func.upper(func.trim(StockBatch.BatchCode)) == batch_dict["BatchCode"]
                    ).first()
                    if db_batch:
                        entry["IsValid"] = False
                        entry["Errors"].append(f"Row {i}: Batch '{batch_dict['BatchCode']}' already exists in inventory")
                    else:
                        entry["InitialBatches"].append(batch_dict)
                else:
                    entry["InitialBatches"].append(batch_dict)

    preview_data = list(groups.values())
    return {"success": True, "data": preview_data, "message": f"Preview generated — {len(preview_data)} medicine(s)"}


# ─────────────────────────────────────────────────────────────────────────────
# Bulk import — creates medicines or adds batches to existing ones atomically
# ─────────────────────────────────────────────────────────────────────────────

@router.post("/import-bulk", summary="Bulk import medicines with optional initial stock")
def bulk_import_medicines(
    medicines: List[MedicineCreate],
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    created = 0
    updated = 0
    stock_sessions: list[str] = []
    try:
        for med_in in medicines:
            initial_stock   = med_in.initial_stock
            existing_med_id = med_in.ExistingMedicineId

            if existing_med_id:
                # ── Medicine already exists — add batches only ─────────────
                existing = db.query(Medicine).filter(Medicine.MedicineId == existing_med_id).first()
                if not existing:
                    raise HTTPException(
                        status_code=404,
                        detail=f"Medicine ID {existing_med_id} not found during import."
                    )
                if initial_stock:
                    ref = _commit_initial_stock_batches(
                        db=db,
                        medicine_id=existing_med_id,
                        batches=initial_stock,
                        user_id=current_user.UserId,
                        medicine_name=existing.BrandName,
                    )
                    stock_sessions.append(ref)
                updated += 1
            else:
                # ── Create new medicine ────────────────────────────────────
                dump = med_in.model_dump(exclude={"initial_stock", "ExistingMedicineId"})
                if not dump.get("Barcode"):
                    dump["Barcode"] = None
                new_med = Medicine(**dump)
                db.add(new_med)
                db.flush()  # get MedicineId before stock creation

                if initial_stock:
                    ref = _commit_initial_stock_batches(
                        db=db,
                        medicine_id=new_med.MedicineId,
                        batches=initial_stock,
                        user_id=current_user.UserId,
                        medicine_name=new_med.BrandName,
                    )
                    stock_sessions.append(ref)
                created += 1

        db.commit()
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Import error: {str(e)}")

    parts = []
    if created: parts.append(f"{created} medicine(s) created")
    if updated: parts.append(f"{updated} existing medicine(s) updated with new batches")
    if stock_sessions: parts.append(f"{len(stock_sessions)} stock session(s) committed")
    return {"success": True, "message": "Import successful — " + ", ".join(parts) if parts else "Nothing to import"}


# ─────────────────────────────────────────────────────────────────────────────
# Update medicine
# ─────────────────────────────────────────────────────────────────────────────

@router.put("/{medicine_id}", response_model=BaseResponse[MedicineResponse], summary="Update a medicine")
def update_medicine(
    medicine_id: int,
    medicine_in: MedicineUpdate,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    medicine = db.query(Medicine).filter(Medicine.MedicineId == medicine_id).first()
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
        
    if medicine_in.Barcode and medicine_in.Barcode != medicine.Barcode:
        existing = db.query(Medicine).filter(Medicine.Barcode == medicine_in.Barcode).first()
        if existing:
            raise HTTPException(status_code=400, detail="Another medicine with this barcode already exists")
            
    update_data = medicine_in.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(medicine, field, value)
        
    db.commit()
    db.refresh(medicine)
    
    response_data = {c.name: getattr(medicine, c.name) for c in medicine.__table__.columns}
    response_data["CategoryName"] = medicine.category.CategoryName if medicine.category else None
    response_data["CompanyName"] = medicine.company.CompanyName if medicine.company else None
    
    return {"data": response_data, "message": "Medicine updated successfully"}


# ─────────────────────────────────────────────────────────────────────────────
# Toggle status
# ─────────────────────────────────────────────────────────────────────────────

@router.put("/{medicine_id}/status", response_model=BaseResponse[MedicineResponse], summary="Toggle medicine status")
def toggle_medicine_status(
    medicine_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    medicine = db.query(Medicine).filter(Medicine.MedicineId == medicine_id).first()
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
        
    medicine.IsActive = not medicine.IsActive
    db.commit()
    db.refresh(medicine)
    return {"data": medicine, "message": f"Medicine status changed to {'Active' if medicine.IsActive else 'Inactive'}"}


# ─────────────────────────────────────────────────────────────────────────────
# Delete medicine
# ─────────────────────────────────────────────────────────────────────────────

@router.delete("/{medicine_id}", summary="Delete a medicine")
def delete_medicine(
    medicine_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    medicine = db.query(Medicine).filter(Medicine.MedicineId == medicine_id).first()
    if not medicine:
        raise HTTPException(status_code=404, detail="Medicine not found")
        
    # Comprehensive Dependency Check
    has_batches = db.query(StockBatch).filter(StockBatch.MedicineId == medicine_id).count() > 0
    has_sales = db.query(SaleItem).filter(SaleItem.BatchId.in_(
        db.query(StockBatch.BatchId).filter(StockBatch.MedicineId == medicine_id)
    )).count() > 0
    has_purchases = db.query(PurchaseItem).filter(PurchaseItem.MedicineId == medicine_id).count() > 0
    
    if has_batches or has_sales or has_purchases:
        logger.warning(f"AUDIT: User {current_user.Username} attempted to delete medicine {medicine_id} ({medicine.BrandName}) but was blocked. Reason: Dependency records exist (Batches: {has_batches}, Sales: {has_sales}, Purchases: {has_purchases}).")
        raise HTTPException(
            status_code=400, 
            detail="Cannot delete medicine because sales, purchases, or stock batch records exist. Please set its status to Inactive instead."
        )
        
    try:
        med_name = medicine.BrandName
        db.delete(medicine)
        db.commit()
        logger.info(f"AUDIT: User {current_user.Username} successfully deleted medicine {medicine_id} ({med_name}).")
        return {"success": True, "message": "Medicine deleted successfully"}
    except IntegrityError as e:
        db.rollback()
        logger.error(f"AUDIT: User {current_user.Username} encountered IntegrityError deleting medicine {medicine_id}. Error: {str(e)}")
        raise HTTPException(
            status_code=400,
            detail="Cannot delete this medicine due to database constraints."
        )
    except Exception as e:
        db.rollback()
        logger.error(f"AUDIT: User {current_user.Username} failed to delete medicine {medicine_id}. Error: {str(e)}")
        raise HTTPException(status_code=500, detail="An error occurred while deleting the medicine")
